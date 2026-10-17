from openpyxl import load_workbook
p = r"C:\Users\LBG\Desktop\New folder\React router\Data\Australia .xlsx"
wb = load_workbook(p, read_only=True, data_only=True)
ws = wb['City_Details']
rows = list(ws.iter_rows(values_only=True))
print('header', rows[0])
for r in rows[1:10]:
    print(r)
