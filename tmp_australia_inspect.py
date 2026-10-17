from openpyxl import load_workbook
p = r"C:\Users\LBG\Desktop\New folder\React router\Data\Australia .xlsx"
wb = load_workbook(p, read_only=True, data_only=True)
print('SHEETS', wb.sheetnames)
for ws in wb.worksheets:
    print('---', ws.title)
    max_rows = min(ws.max_row, 12)
    for row in ws.iter_rows(min_row=1, max_row=max_rows, values_only=True):
        print(row)
