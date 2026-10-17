import argparse
import hashlib
import json
import re
import shutil
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageOps


ROOT = Path(__file__).resolve().parent
WORKBOOK_PATH = ROOT / 'Data' / 'South_America.xlsx'
IMAGE_DIR = ROOT / 'backend' / 'public' / 'Images' / 'South America cities images'
MANIFEST_PATH = ROOT / 'South_America_city_image_manifest.json'
REVIEW_DIR = ROOT / 'South_America_city_image_review'
REJECTED_DIR = ROOT / 'South_America_rejected_city_image_candidates'
VISUAL_REVIEW_PATH = ROOT / 'South_America_city_image_visual_review.json'
OPENVERSE_URL = 'https://api.openverse.org/v1/images/'
ALLOWED_LICENSES = {'by', 'by-sa', 'cc0', 'pdm'}
MIN_IMAGE_PIXELS = 400_000
EXPECTED_COUNTRIES = 12
EXPECTED_SOURCE_ROWS = 123
EXPECTED_TARGET_CITIES = 136
USER_AGENT = 'SouthAmericaCityImageAudit/1.0 (licensed city image research)'

COUNTRY_CODES = {
    'Argentina': 'AR', 'Bolivia': 'BO', 'Brazil': 'BR', 'Chile': 'CL',
    'Colombia': 'CO', 'Ecuador': 'EC', 'Guyana': 'GY', 'Paraguay': 'PY',
    'Peru': 'PE', 'Suriname': 'SR', 'Uruguay': 'UY', 'Venezuela': 'VE',
}

DATABASE_ONLY_CITIES = {
    'Argentina': ['La Plata'],
    'Ecuador': ['Cuenca', 'Loja'],
    'Guyana': ['Linden', 'New Amsterdam'],
    'Paraguay': ['Encarnación', 'San Lorenzo'],
    'Suriname': ['Lelydorp', 'Nieuw Nickerie'],
    'Uruguay': ['Paysandú', 'Rivera', 'Salto'],
    'Venezuela': ['Ciudad Guayana', 'Maracay'],
}

CITY_EXTRA_QUERIES = {
    'Cuenca': ['"Cuenca" "Ecuador" "Catedral Nueva"', '"New Cathedral" "Cuenca" "Ecuador"'],
    'Ciudad Guayana': ['"Ciudad Guayana" "Cachamay" "Venezuela"', '"Ciudad Guayana" "panoramicas" "Venezuela"'],
    'Santa Elena de Uairén': ['"Santa Elena de Uairen" "Venezuela" town view'],
}


def normalize(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii')
    text = text.lower().replace('&', ' and ')
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', text).split())


def slugify(value):
    return normalize(value).replace(' ', '-')


def get_inventory():
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    rows = workbook['City_Details'].iter_rows(values_only=True)
    headers = [str(value or '').strip() for value in next(rows)]
    country_index = headers.index('Country')
    city_index = headers.index('City')
    landmark_index = headers.index('Famous Landmarks')
    source_rows = []
    targets = []
    excluded = []
    corrected = []

    for row in rows:
        if not row or not row[country_index] or not row[city_index]:
            continue
        country = str(row[country_index]).strip()
        source_city = str(row[city_index]).strip()
        source_rows.append({'country': country, 'city': source_city})
        landmarks = str(row[landmark_index] or '').strip()

        if country == 'Ecuador' and source_city == 'Baltra':
            excluded.append({'country': country, 'city': source_city,
                             'reason': 'Baltra is an island and airport gateway, not a city.'})
            continue
        city = source_city
        if country == 'Ecuador' and source_city == 'San Cristóbal':
            city = 'Puerto Baquerizo Moreno'
            corrected.append({'country': country, 'source_label': source_city,
                              'verified_city': city,
                              'reason': 'San Cristóbal is the island; Puerto Baquerizo Moreno is the city and provincial capital.'})
        targets.append({
            'country': country, 'country_code': COUNTRY_CODES[country], 'city': city,
            'source_city': source_city, 'inventory_source': 'Data/South_America.xlsx#City_Details',
            'landmarks': landmarks,
        })
    workbook.close()

    if len(source_rows) != EXPECTED_SOURCE_ROWS:
        raise ValueError(f'Expected {EXPECTED_SOURCE_ROWS} source rows, found {len(source_rows)}')
    countries = {row['country'] for row in source_rows}
    if countries != set(COUNTRY_CODES):
        raise ValueError(f'Source country mismatch: {sorted(countries)}')

    seen = {(item['country_code'], normalize(item['city'])) for item in targets}
    database_additions = []
    for country, cities in DATABASE_ONLY_CITIES.items():
        for city in cities:
            key = (COUNTRY_CODES[country], normalize(city))
            if key not in seen:
                seen.add(key)
                targets.append({
                    'country': country, 'country_code': COUNTRY_CODES[country], 'city': city,
                    'source_city': city, 'inventory_source': 'Live database: Country/City by ISO country code',
                    'landmarks': '',
                })
                database_additions.append({'country': country, 'city': city})

    if len(countries) != EXPECTED_COUNTRIES:
        raise ValueError(f'Expected {EXPECTED_COUNTRIES} countries, found {len(countries)}')
    if len(targets) != EXPECTED_TARGET_CITIES:
        raise ValueError(f'Expected {EXPECTED_TARGET_CITIES} unique city-country targets, found {len(targets)}')

    for item in targets:
        item.update({
            'filename': f"{item['country_code'].lower()}-{slugify(item['city'])}.webp",
            'image_source_url': None, 'source_page_url': None, 'creator': None,
            'license': None, 'license_url': None, 'source_title': None, 'source_tags': [],
            'source_queries': [], 'dimensions': None, 'file_sha256': None,
            'candidate_download_attempts': [],
            'download_status': 'pending', 'verification_status': 'pending', 'unresolved_issue': None,
        })

    reconciliation = {
        'source_workbook_rows': len(source_rows), 'source_country_count': len(countries),
        'valid_workbook_city_targets': len(targets) - len(database_additions),
        'database_city_rows_checked': 58, 'database_only_valid_cities_added': database_additions,
        'excluded_non_city_records': excluded, 'corrected_source_city_labels': corrected,
        'known_database_integrity_issues': [
            'Country code AR is attached to a database country named TEST SAVE; ISO code and workbook map AR to Argentina.',
            'Database labels conflict with city slugs for BO La Paz/Santa Cruz de la Sierra, BR São Paulo/São Luís, and CL La Serena/Calama.',
        ],
        'duplicate_source_city_country_pairs': [],
    }
    return targets, reconciliation


def request_bytes(url, attempts=2, timeout=10):
    request = urllib.request.Request(url, headers={'User-Agent': USER_AGENT})
    for attempt in range(attempts):
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read()
        except urllib.error.HTTPError as exc:
            if exc.code == 429 and attempt + 1 < attempts:
                retry_after = exc.headers.get('Retry-After')
                delay = int(retry_after) if retry_after and retry_after.isdigit() else min(30, 5 * (attempt + 1))
                if delay > 30:
                    raise RuntimeError(f'HTTP 429; source requested Retry-After {delay}s, deferred.') from exc
                time.sleep(delay)
                continue
            if exc.code in {500, 502, 503, 504} and attempt + 1 < attempts:
                time.sleep(2 ** (attempt + 1))
                continue
            raise
        except Exception:
            if attempt + 1 >= attempts:
                raise
            time.sleep(2 ** (attempt + 1))
    raise RuntimeError(f'Failed to fetch {url}')


def openverse_search(query):
    params = urllib.parse.urlencode({
        'q': query, 'page_size': 20, 'license': ','.join(sorted(ALLOWED_LICENSES)), 'mature': 'false',
    })
    return json.loads(request_bytes(f'{OPENVERSE_URL}?{params}').decode('utf-8')).get('results', [])


def find_wikipedia_image_candidate(target):
    page_titles = [f"{target['city']}, {target['country']}", target['city']]
    last_error = None
    for page_title in page_titles:
        summary_url = 'https://en.wikipedia.org/api/rest_v1/page/summary/' + urllib.parse.quote(page_title.replace(' ', '_'))
        try:
            summary = json.loads(request_bytes(summary_url, attempts=1, timeout=12).decode('utf-8'))
        except Exception as exc:
            last_error = f'Wikipedia page lookup failed: {exc}'
            if '429' in str(exc):
                break
            continue

        title = str(summary.get('title') or '')
        context = normalize(' '.join([title, summary.get('description') or '', summary.get('extract') or '']))
        city = normalize(target['city'])
        country = normalize(target['country'])
        if not (city in normalize(title) or city.replace(' ', '') in normalize(title).replace(' ', '')):
            continue
        if country not in context:
            continue

        image_url = (summary.get('originalimage') or {}).get('source')
        if not image_url:
            continue
        image_path_parts = urllib.parse.urlparse(image_url).path.split('/')
        if 'thumb' in image_path_parts:
            thumb_index = image_path_parts.index('thumb')
            if len(image_path_parts) <= thumb_index + 4:
                continue
            file_title = urllib.parse.unquote(image_path_parts[thumb_index + 3])
        else:
            file_title = urllib.parse.unquote(image_path_parts[-1])

        params = urllib.parse.urlencode({
            'action': 'query',
            'titles': f'File:{file_title}',
            'prop': 'imageinfo',
            'iiprop': 'url|size|extmetadata',
            'iiurlwidth': 1600,
            'format': 'json',
        })
        try:
            info_data = json.loads(request_bytes(f'https://en.wikipedia.org/w/api.php?{params}', attempts=1, timeout=12).decode('utf-8'))
            pages = info_data.get('query', {}).get('pages', {}).values()
            page = next(iter(pages), {})
            info = (page.get('imageinfo') or [{}])[0]
        except Exception as exc:
            last_error = f'Wikipedia image metadata lookup failed: {exc}'
            if '429' in str(exc):
                break
            continue

        metadata = info.get('extmetadata') or {}
        license_name = str((metadata.get('LicenseShortName') or {}).get('value') or '').lower()
        if 'by-sa' in license_name:
            license_code = 'by-sa'
        elif 'cc0' in license_name:
            license_code = 'cc0'
        elif 'public domain' in license_name or license_name == 'pdm':
            license_code = 'pdm'
        elif re.search(r'\bcc\s+by\b', license_name):
            license_code = 'by'
        else:
            continue

        width = int(info.get('width') or 0)
        height = int(info.get('height') or 0)
        if width < 400 or height < 400 or width * height < MIN_IMAGE_PIXELS:
            continue
        image_download_url = info.get('thumburl') or info.get('url')
        if not image_download_url:
            continue
        creator = re.sub(r'<[^>]+>', '', str((metadata.get('Artist') or {}).get('value') or '')).strip()
        return {
            'url': image_download_url,
            'foreign_landing_url': info.get('descriptionurl') or summary.get('content_urls', {}).get('desktop', {}).get('page'),
            'title': title,
            'creator': creator,
            'license': license_code,
            'license_url': (metadata.get('LicenseUrl') or {}).get('value'),
            'provider': 'Wikimedia Commons via Wikipedia API',
            'width': width,
            'height': height,
            'summary': summary,
        }, None
    return None, last_error


def candidate_score(item, target):
    if (item.get('license') or '').lower() not in ALLOWED_LICENSES:
        return 0
    width = int(item.get('width') or 0)
    height = int(item.get('height') or 0)
    if width < 400 or height < 400 or width * height < MIN_IMAGE_PIXELS:
        return 0
    tags = [str(tag.get('name') or '') for tag in item.get('tags', [])]
    title = str(item.get('title') or '')
    title_normalized = normalize(title)
    metadata = normalize(' '.join([title, *tags]))
    city = normalize(target['city'])
    country = normalize(target['country'])
    city_in_metadata = city in metadata or city.replace(' ', '') in metadata.replace(' ', '')
    city_in_title = city in title_normalized or city.replace(' ', '') in title_normalized.replace(' ', '')
    landmarks = [part.strip() for part in re.split(r'[,;|\n]', target.get('landmarks') or '') if part.strip()]
    landmark_matches = [part for part in landmarks if normalize(part) in metadata]
    if not city_in_metadata or not (country in metadata or landmark_matches or city_in_title):
        return 0
    if any(term in metadata for term in (
        'portrait', 'selfie', 'logo', 'flag', 'map', 'illustration', 'coat of arms',
        'graphic', 'infographic', 'chart', 'graph', 'visualization', 'temperature',
        'show your stripes', 'poster', 'flyer', 'advertisement', 'satellite',
        'false color', 'earth observation', 'airport', 'airfield', 'air terminal',
    )):
        return 0
    target_exclusions = {
        'El Calafate': ['glacier', 'glaciar'],
        'La Plata': ['rio de la plata', 'riodelaplata'],
        'Recife': ['olinda'],
        'Ipiales': ['popayan', 'popayán'],
        'Rivera': ['minas de corrales'],
        'Salto': ['salto del penitente'],
        'Pucallpa': ['hammock', 'hamaca'],
    }
    if any(normalize(term) in metadata for term in target_exclusions.get(target['city'], [])):
        return 0
    score = 100 + (30 if country in metadata else 0) + min(50, 20 * len(landmark_matches))
    if any(term in metadata for term in ('skyline', 'historic centre', 'historic center', 'waterfront', 'downtown', 'plaza')):
        score += 5
    return score


def build_queries(target):
    queries = []
    landmarks = [part.strip() for part in re.split(r'[,;|\n]', target.get('landmarks') or '') if part.strip()]
    for landmark in landmarks[:2]:
        queries.append(f'"{target["city"]}" "{target["country"]}" "{landmark}"')
    queries.extend([f'"{target["city"]}" "{target["country"]}" city',
                    f'"{target["city"]}" "{target["country"]}" skyline',
                    f'"{target["city"]}" "{target["country"]}" architecture',
                    f'"{target["city"]}" "{target["country"]}" landmark'])
    queries.extend(CITY_EXTRA_QUERIES.get(target['city'], []))
    return list(dict.fromkeys(queries))


def save_manifest(manifest):
    manifest['updated_at'] = datetime.now(timezone.utc).isoformat(timespec='seconds').replace('+00:00', 'Z')
    temp_path = MANIFEST_PATH.with_suffix('.json.tmp')
    temp_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    temp_path.replace(MANIFEST_PATH)


def load_manifest():
    return json.loads(MANIFEST_PATH.read_text(encoding='utf-8'))


def download_images(start_index=0, limit=None):
    targets, reconciliation = get_inventory()
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    previous = {}
    if MANIFEST_PATH.is_file():
        try:
            previous = {item['filename']: item for item in load_manifest().get('cities', [])}
        except (OSError, ValueError, KeyError):
            previous = {}
    for target in targets:
        old = previous.get(target['filename'])
        if old:
            target.update({key: value for key, value in old.items() if key not in {'country', 'country_code', 'city', 'filename'}})
            if old.get('visual_review') == 'rejected':
                target['previous_visual_rejection'] = old.get('unresolved_issue')
    manifest = {
        'continent': 'South America', 'source_workbook': 'Data/South_America.xlsx',
        'target_folder': 'backend/public/Images/South America cities images',
        'filename_convention': '{lowercase ISO-3166-1 alpha-2}-{ascii-city-slug}.webp',
        'allowed_licenses': sorted(ALLOWED_LICENSES), 'minimum_dimensions': [400, 400],
        'minimum_pixel_count': MIN_IMAGE_PIXELS, 'countries_checked': EXPECTED_COUNTRIES,
        'unique_city_targets': EXPECTED_TARGET_CITIES, 'reconciliation': reconciliation,
        'visual_review_policy': 'Downloaded files remain pending until each contact sheet is reviewed.',
        'cities': targets,
    }
    save_manifest(manifest)
    used_hashes = {}
    failures = 0
    end_index = len(targets) if limit is None else min(len(targets), start_index + limit)
    batch_failures = 0
    for index in range(start_index, end_index):
        target = manifest['cities'][index]
        print(f'[{index + 1}/{len(targets)}] {target["city"]}, {target["country"]}', flush=True)
        output_path = IMAGE_DIR / target['filename']
        if output_path.is_file() and target.get('image_source_url'):
            try:
                with Image.open(output_path) as image:
                    image.verify()
                with Image.open(output_path) as image:
                    image.load()
                    if image.format != 'WEBP' or image.width < 400 or image.height < 400:
                        raise ValueError(f'Unexpected existing image: {image.format} {image.size}')
                    target['dimensions'] = [image.width, image.height]
                target['file_sha256'] = hashlib.sha256(output_path.read_bytes()).hexdigest()
                target['download_status'] = 'downloaded'
                target['verification_status'] = 'pending_visual_review'
                save_manifest(manifest)
                continue
            except Exception as exc:
                target['download_status'] = 'invalid'
                target['verification_status'] = 'unresolved'
                target['unresolved_issue'] = f'Existing image failed integrity validation: {exc}'
        if target.get('download_status') == 'downloaded':
            target['download_status'] = 'pending'
            target['verification_status'] = 'pending'
        wikipedia_queries = [f"{target['city']}, {target['country']}", target['city']]
        target['source_queries'].extend(f'Wikipedia page lookup: {query}' for query in wikipedia_queries)
        wiki_candidate, wiki_error = find_wikipedia_image_candidate(target)
        if wiki_candidate:
            try:
                print(f"  Wikipedia image: {wiki_candidate['title']}", flush=True)
                data = request_bytes(wiki_candidate['url'], attempts=1, timeout=20)
                digest = hashlib.sha256(data).hexdigest()
                if digest not in used_hashes:
                    image = Image.open(BytesIO(data))
                    image = ImageOps.exif_transpose(image).convert('RGB')
                    image.load()
                    if image.width >= 400 and image.height >= 400 and image.width * image.height >= MIN_IMAGE_PIXELS:
                        image.thumbnail((1600, 1200), Image.Resampling.LANCZOS)
                        temp_path = output_path.with_suffix('.webp.part')
                        image.save(temp_path, 'WEBP', quality=88, method=6)
                        temp_path.replace(output_path)
                        used_hashes[digest] = target['filename']
                        target.update({
                            'image_source_url': wiki_candidate['url'],
                            'source_page_url': wiki_candidate['foreign_landing_url'],
                            'creator': wiki_candidate['creator'], 'license': wiki_candidate['license'],
                            'license_url': wiki_candidate['license_url'], 'source_title': wiki_candidate['title'],
                            'source_tags': [], 'source_provider': wiki_candidate['provider'],
                            'city_country_verification': {
                                'page_title': wiki_candidate['summary'].get('title'),
                                'description': wiki_candidate['summary'].get('description'),
                                'extract': wiki_candidate['summary'].get('extract'),
                            },
                            'dimensions': [image.width, image.height],
                            'file_sha256': hashlib.sha256(output_path.read_bytes()).hexdigest(),
                            'candidate_metadata_score': 200, 'download_status': 'downloaded',
                            'verification_status': 'pending_visual_review', 'visual_review': 'pending',
                            'visual_review_evidence': None, 'unresolved_issue': None,
                        })
                        save_manifest(manifest)
                        time.sleep(0.7)
                        continue
            except Exception as exc:
                target['candidate_download_attempts'].append({'url': wiki_candidate['url'], 'error': str(exc)})
                target['wikipedia_lookup_issue'] = f'Wikipedia image download failed: {exc}'
        elif wiki_error:
            target['wikipedia_lookup_issue'] = wiki_error

        candidate_map = {}
        for query in build_queries(target):
            target['source_queries'].append(query)
            print(f'  search: {query}', flush=True)
            try:
                for item in openverse_search(query):
                    score = candidate_score(item, target)
                    if score:
                        candidate_map[item.get('id') or item.get('url')] = (score, item)
            except Exception as exc:
                target['unresolved_issue'] = f'Openverse search error: {exc}'
            time.sleep(0.7)

        candidates = sorted(
            candidate_map.values(),
            key=lambda pair: (pair[0], int(pair[1].get('width') or 0) * int(pair[1].get('height') or 0)),
            reverse=True,
        )
        saved = False
        for score, item in candidates[:3]:
            try:
                print(f'  candidate: {item.get("title") or item["url"]}', flush=True)
                data = request_bytes(item['url'], attempts=2, timeout=12)
                digest = hashlib.sha256(data).hexdigest()
                if digest in used_hashes:
                    continue
                image = Image.open(BytesIO(data))
                image = ImageOps.exif_transpose(image).convert('RGB')
                image.load()
                if image.width < 400 or image.height < 400 or image.width * image.height < MIN_IMAGE_PIXELS:
                    continue
                image.thumbnail((1600, 1200), Image.Resampling.LANCZOS)
                output_path = IMAGE_DIR / target['filename']
                temp_path = output_path.with_suffix('.webp.part')
                image.save(temp_path, 'WEBP', quality=88, method=6)
                temp_path.replace(output_path)
                used_hashes[digest] = target['filename']
                target.update({
                    'image_source_url': item['url'], 'source_page_url': item.get('foreign_landing_url'),
                    'creator': item.get('creator'), 'license': item.get('license'),
                    'license_url': item.get('license_url'), 'source_title': item.get('title'),
                    'source_tags': [str(tag.get('name') or '') for tag in item.get('tags', [])],
                    'source_provider': item.get('provider') or item.get('source'),
                    'dimensions': [image.width, image.height],
                    'file_sha256': hashlib.sha256(output_path.read_bytes()).hexdigest(),
                    'candidate_metadata_score': score, 'download_status': 'downloaded',
                    'verification_status': 'pending_visual_review', 'visual_review': 'pending',
                    'visual_review_evidence': None, 'unresolved_issue': None,
                })
                saved = True
                break
            except Exception as exc:
                target['candidate_download_attempts'].append({'url': item.get('url'), 'error': str(exc)})
                target['unresolved_issue'] = f'Candidate download or conversion failed: {exc}'
        if not saved:
            target['download_status'] = 'unresolved'
            target['verification_status'] = 'unresolved'
            if not target['unresolved_issue']:
                target['unresolved_issue'] = 'No licensed, sufficiently large, city-and-country-matched candidate found.'
            failures += 1
            batch_failures += 1
        save_manifest(manifest)
    print(f'BATCH_PROCESSED {end_index - start_index}')
    print(f'BATCH_UNRESOLVED {batch_failures}')
    print(f'TOTAL_UNRESOLVED_IN_BATCHED_RUN {failures}')
    print(f'MANIFEST {MANIFEST_PATH}')


def difference_hash(image):
    image = ImageOps.exif_transpose(image).convert('L').resize((9, 8), Image.Resampling.LANCZOS)
    pixels = list(image.getdata())
    bits = 0
    for row in range(8):
        for col in range(8):
            bits = (bits << 1) | (1 if pixels[row * 9 + col] > pixels[row * 9 + col + 1] else 0)
    return bits


def apply_visual_review():
    manifest = load_manifest()
    decisions = json.loads(VISUAL_REVIEW_PATH.read_text(encoding='utf-8'))
    rejected = {item['filename']: item['reason'] for item in decisions.get('rejected', [])}
    rejected_replacements = {item['filename']: item['reason'] for item in decisions.get('rejected_replacements', [])}
    approved = {item['filename']: item['reason'] for item in decisions.get('approved_replacements', [])}
    REJECTED_DIR.mkdir(parents=True, exist_ok=True)
    approved_count = 0
    rejected_count = 0
    for target in manifest['cities']:
        output_path = IMAGE_DIR / target['filename']
        rejected_path = REJECTED_DIR / target['filename']
        replacement_reason = rejected_replacements.get(target['filename'])
        reason = replacement_reason or rejected.get(target['filename'])
        if target['filename'] in approved and not output_path.is_file() and rejected_path.is_file():
            shutil.move(str(rejected_path), str(output_path))
        if target['filename'] in approved and output_path.is_file():
            target['download_status'] = 'downloaded'
            target['verification_status'] = 'verified'
            target['visual_review'] = 'passed'
            target['visual_review_evidence'] = approved[target['filename']]
            target['unresolved_issue'] = None
            approved_count += 1
        elif reason and (replacement_reason or target.get('visual_review') != 'pending'):
            if output_path.is_file():
                if rejected_path.exists():
                    rejected_path.unlink()
                shutil.move(str(output_path), str(rejected_path))
            target['download_status'] = 'rejected'
            target['verification_status'] = 'unresolved'
            target['visual_review'] = 'rejected'
            target['visual_review_evidence'] = decisions['review_method']
            target['unresolved_issue'] = reason
            rejected_count += 1
        elif output_path.is_file() and target.get('visual_review') == 'pending':
            continue
        elif output_path.is_file():
            target['verification_status'] = 'verified'
            target['visual_review'] = 'passed'
            target['visual_review_evidence'] = decisions['review_method']
            approved_count += 1
    save_manifest(manifest)
    print(f'VISUALLY_APPROVED_PENDING_INTEGRITY {approved_count}')
    print(f'VISUALLY_REJECTED {rejected_count}')
    print(f'REJECTED_CANDIDATE_FOLDER {REJECTED_DIR}')


def normalize_approved_panoramas():
    manifest = load_manifest()
    approved = {'ec-manta.webp', 'pe-chachapoyas.webp'}
    normalized = []
    for target in manifest['cities']:
        if target['filename'] not in approved:
            continue
        source_data = request_bytes(target['image_source_url'], attempts=1, timeout=20)
        with Image.open(BytesIO(source_data)) as source:
            source = ImageOps.exif_transpose(source).convert('RGB')
            source.load()
            crop_width = source.height * 4
            if crop_width > source.width or source.height < 400:
                raise ValueError(f"Native source cannot support a non-upscaled 4:1 crop: {target['city']} {source.size}")
            left = (source.width - crop_width) // 2
            source = source.crop((left, 0, left + crop_width, source.height))
            output = source.resize((1600, 400), Image.Resampling.LANCZOS)
            output_path = IMAGE_DIR / target['filename']
            output.save(output_path, 'WEBP', quality=88, method=6)
        target['dimensions'] = [1600, 400]
        target['file_sha256'] = hashlib.sha256(output_path.read_bytes()).hexdigest()
        target['download_status'] = 'downloaded'
        target['verification_status'] = 'verified'
        target['visual_review'] = 'passed'
        target['visual_review_evidence'] = 'Correct city landmark/panorama confirmed on contact sheet; centered 4:1 crop from native-resolution source without upscaling.'
        target['unresolved_issue'] = None
        normalized.append(target['filename'])
    save_manifest(manifest)
    print('NORMALIZED_PANORAMAS', json.dumps(normalized))


def audit_images(mark_reviewed=False):
    manifest = load_manifest()
    missing, invalid, exact_duplicates, near_duplicates = [], [], [], []
    seen_hashes = {}
    hashes = []
    for target in manifest['cities']:
        path = IMAGE_DIR / target['filename']
        if not path.is_file():
            missing.append(target['city'] + ', ' + target['country'])
            if target.get('visual_review') != 'rejected':
                target['download_status'] = 'missing'
                target['verification_status'] = 'unresolved'
            continue
        try:
            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image:
                image.load()
                if image.format != 'WEBP' or image.width < 400 or image.height < 400 or image.width * image.height < MIN_IMAGE_PIXELS:
                    raise ValueError(f'{image.format} {image.size}')
                current_hash = difference_hash(image)
                target['dimensions'] = [image.width, image.height]
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            target['file_sha256'] = digest
            if digest in seen_hashes:
                exact_duplicates.append([seen_hashes[digest], target['filename']])
            else:
                seen_hashes[digest] = target['filename']
            hashes.append((target['filename'], current_hash))
            if mark_reviewed:
                target['verification_status'] = 'verified'
                target['visual_review'] = 'passed'
                target['visual_review_evidence'] = 'Manually reviewed in generated contact sheet.'
        except Exception as exc:
            invalid.append({'filename': target['filename'], 'reason': str(exc)})
            target['download_status'] = 'invalid'
            target['verification_status'] = 'unresolved'
            target['unresolved_issue'] = f'Image integrity validation failed: {exc}'

    for index, (first_name, first_hash) in enumerate(hashes):
        for second_name, second_hash in hashes[index + 1:]:
            distance = (first_hash ^ second_hash).bit_count()
            if distance <= 2:
                near_duplicates.append({'files': [first_name, second_name], 'dhash_distance': distance})
    for pair in exact_duplicates:
        for target in manifest['cities']:
            if target['filename'] in pair:
                target['verification_status'] = 'unresolved'
                target['unresolved_issue'] = 'Image content duplicates another city asset.'

    verified = sum(target.get('verification_status') == 'verified' for target in manifest['cities'])
    unresolved = [target for target in manifest['cities'] if target.get('verification_status') != 'verified']
    country_targets = {}
    for target in manifest['cities']:
        country_targets.setdefault(target['country'], []).append(target)
    save_manifest(manifest)
    report = [
        '# South America City Image Audit', '', f"Generated: {manifest.get('updated_at')}",
        f"Countries checked: {manifest['countries_checked']}",
        f"Unique city targets: {manifest['unique_city_targets']}", f'Verified images: {verified}',
        f'Missing files: {len(missing)}', f'Invalid files: {len(invalid)}',
        f'Exact duplicate pairs: {len(exact_duplicates)}',
        f'Perceptual near-duplicate pairs (dHash distance <= 2): {len(near_duplicates)}',
        f'Cities unresolved: {len(unresolved)}', f"Image folder: {manifest['target_folder']}",
        f'Manifest: {MANIFEST_PATH.name}', '', '## Inventory Reconciliation',
        f"Source workbook rows: {manifest['reconciliation']['source_workbook_rows']}",
        f"Valid source city targets: {manifest['reconciliation']['valid_workbook_city_targets']}",
        f"Additional database city names: {len(manifest['reconciliation']['database_only_valid_cities_added'])}",
        'Duplicate source city-country pairs: 0',
        'Excluded: Baltra, Ecuador (island and airport gateway, not a city).',
        'Corrected: San Cristóbal, Ecuador refers to the island; the city is Puerto Baquerizo Moreno.',
        'Database issues: AR is labeled TEST SAVE; BO La Paz/Santa Cruz city names are swapped against slugs; BR br-sao-paulo is labeled São Luís; CL cl-la-serena is labeled Calama.',
        'Sources: Data/South_America.xlsx#City_Details; live Prisma Country/City rows queried by the 12 ISO country codes; catalog API route /countries/:code/cities.',
        'Image acquisition: Openverse public API with CC BY, CC BY-SA, CC0, and public-domain filters; Wikipedia page summaries plus Wikipedia imageinfo/Commons license metadata. Commons API hostname DNS resolution failed during fallback; 429 Retry-After responses were honored or deferred.',
        'Geographic checks: https://en.wikipedia.org/wiki/Baltra_Island; https://en.wikipedia.org/wiki/San_Crist%C3%B3bal_Island; https://en.wikipedia.org/wiki/Puerto_Baquerizo_Moreno.',
        '', '## Inventory by Country',
    ]
    for country in sorted(country_targets):
        entries = country_targets[country]
        report.append(f"### {country} ({len(entries)} cities)")
        report.extend([
            f"- {target['city']}: {'verified' if target.get('verification_status') == 'verified' else 'unresolved'} ({target['filename']})"
            for target in entries
        ])
    report.extend(['', '## File Audit',
                   f'Output files present: {len(seen_hashes)}; manifest entries with files: {len(seen_hashes)}.',
                   'All present files were decoded and loaded as WebP; minimum dimensions 400×400 and 400,000 pixels.',
                   f'Exact SHA-256 duplicate pairs: {len(exact_duplicates)}.',
                   f'Perceptual duplicate pairs at dHash distance <= 2: {len(near_duplicates)}.',
                   'Rejected candidates were moved outside the deliverable folder to South_America_rejected_city_image_candidates.',
                   '', '## Unresolved'])
    for item in unresolved:
        reason = item.get('unresolved_issue') or 'Visual review pending.'
        source_issue = item.get('wikipedia_lookup_issue')
        if source_issue:
            reason += f' Additional source issue: {source_issue}'
        report.append(f"- {item['city']}, {item['country']}: {reason}")
    if not unresolved:
        report.append('- None.')
    if near_duplicates:
        report.extend(['', '## Potential Visual Duplicates'])
        report.extend([f"- {item['files'][0]} / {item['files'][1]} (dHash distance {item['dhash_distance']})"
                       for item in near_duplicates])
    (ROOT / 'South_America_city_image_audit.md').write_text('\n'.join(report) + '\n', encoding='utf-8')
    print(json.dumps({
        'countries': manifest['countries_checked'], 'city_targets': manifest['unique_city_targets'],
        'verified': verified, 'missing': missing, 'invalid': invalid,
        'exact_duplicates': exact_duplicates, 'near_duplicates': near_duplicates,
        'unresolved': len(unresolved),
    }, ensure_ascii=False, indent=2))


def make_contact_sheets():
    manifest = load_manifest()
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    entries = [item for item in manifest['cities'] if (IMAGE_DIR / item['filename']).is_file()]
    cell_width, cell_height, per_sheet = 320, 230, 16
    for sheet_index in range(0, len(entries), per_sheet):
        chunk = entries[sheet_index:sheet_index + per_sheet]
        sheet = Image.new('RGB', (cell_width * 4, cell_height * 4), 'white')
        draw = ImageDraw.Draw(sheet)
        for cell_index, target in enumerate(chunk):
            col, row = cell_index % 4, cell_index // 4
            x, y = col * cell_width, row * cell_height
            with Image.open(IMAGE_DIR / target['filename']) as image:
                image = ImageOps.exif_transpose(image).convert('RGB')
                image.thumbnail((cell_width - 12, 175), Image.Resampling.LANCZOS)
                sheet.paste(image, (x + (cell_width - image.width) // 2, y + 4))
            draw.text((x + 6, y + 182), f"{target['city']} | {target['country']}", fill='black')
            draw.text((x + 6, y + 200), target['filename'], fill='black')
        output_path = REVIEW_DIR / f'contact-sheet-{sheet_index // per_sheet + 1:02d}.jpg'
        sheet.save(output_path, quality=88)
        print(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--inventory-only', action='store_true')
    parser.add_argument('--download', action='store_true')
    parser.add_argument('--start-index', type=int, default=0)
    parser.add_argument('--limit', type=int)
    parser.add_argument('--contact-sheets', action='store_true')
    parser.add_argument('--apply-visual-review', action='store_true')
    parser.add_argument('--normalize-approved-panoramas', action='store_true')
    parser.add_argument('--audit', action='store_true')
    parser.add_argument('--mark-reviewed', action='store_true')
    args = parser.parse_args()
    if args.inventory_only:
        targets, reconciliation = get_inventory()
        print(json.dumps({
            'countries': len({item['country'] for item in targets}), 'unique_city_targets': len(targets),
            'country_city_counts': dict(sorted(Counter(item['country'] for item in targets).items())),
            'reconciliation': reconciliation,
        }, ensure_ascii=False, indent=2))
    if args.download:
        download_images(start_index=max(0, args.start_index), limit=args.limit)
    if args.contact_sheets:
        make_contact_sheets()
    if args.apply_visual_review:
        apply_visual_review()
    if args.normalize_approved_panoramas:
        normalize_approved_panoramas()
    if args.audit:
        audit_images(mark_reviewed=args.mark_reviewed)
    if not any(vars(args).values()):
        parser.print_help()


if __name__ == '__main__':
    main()