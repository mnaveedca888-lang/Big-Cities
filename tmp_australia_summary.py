from openpyxl import load_workbook
from collections import Counter
p = r"C:\Users\LBG\Desktop\New folder\React router\Data\Australia .xlsx"
wb = load_workbook(p, read_only=True, data_only=True)
country_ws = wb['Country_Details']
city_ws = wb['City_Details']
country_rows = list(country_ws.iter_rows(values_only=True))
city_rows = list(city_ws.iter_rows(values_only=True))
print('country rows total', len(country_rows)-1)
print('city rows total', len(city_rows)-1)
headers = country_rows[0]
for i,h in enumerate(headers):
    print(i, h)
print('country names sample', [r[0] for r in country_rows[1:11]])
print('continent values', Counter(r[3] for r in country_rows[1:]))
print('city country sample', [r[0] for r in city_rows[1:20]])
print('city country count', Counter(r[0] for r in city_rows[1:]))
