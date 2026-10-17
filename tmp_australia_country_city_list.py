from openpyxl import load_workbook
from collections import defaultdict
import json

p = r"C:\Users\LBG\Desktop\New folder\React router\Data\Australia .xlsx"
wb = load_workbook(p, read_only=True, data_only=True)
country_rows = list(wb['Country_Details'].iter_rows(values_only=True))
city_rows = list(wb['City_Details'].iter_rows(values_only=True))

country_headers = country_rows[0]
country_data = []
for row in country_rows[1:]:
    item = dict(zip(country_headers, row))
    country = (item.get('country') or '').strip()
    if country:
        country_data.append({'country': country, 'capital': item.get('Capital'), 'population': item.get('Population'), 'continent': item.get('Continent')})

city_headers = city_rows[0]
city_data = []
for row in city_rows[1:]:
    item = dict(zip(city_headers, row))
    country = (item.get('Country Name') or '').strip()
    city = (item.get('City') or '').strip()
    if country and city:
        city_data.append({'country': country, 'city': city})

city_by_country = defaultdict(list)
for item in city_data:
    city_by_country[item['country']].append(item['city'])

unique_countries = []
seen = set()
for item in country_data:
    country = item['country']
    if country not in seen:
        seen.add(country)
        unique_countries.append(country)

result = {
    'continent_label': 'Oceania',
    'countries': unique_countries,
    'total_countries': len(unique_countries),
    'total_cities': len(city_data),
    'cities_by_country': {country: city_by_country.get(country, []) for country in unique_countries},
    'city_list': city_data,
}

out = r"C:\Users\LBG\Desktop\New folder\React router\tmp_australia_country_city_list.json"
with open(out, 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)
print('WROTE', out)
print('TOTAL_COUNTRIES', len(unique_countries))
print('TOTAL_CITIES', len(city_data))
