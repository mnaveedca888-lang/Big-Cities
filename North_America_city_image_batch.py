import json
import re
import shutil
import time
import unicodedata
import urllib.parse
import urllib.request
import urllib.error
import sys
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageOps
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
WORKBOOK_PATH = ROOT / 'Data' / 'North America.xlsx'
OUTDIR = ROOT / 'backend' / 'public' / 'Images' / 'North America cities images'
OUTDIR.mkdir(parents=True, exist_ok=True)
MANIFEST_PATH = ROOT / 'North_America_city_image_manifest.json'
QUARANTINE_DIR = ROOT / 'North_America_rejected_city_image_candidates'
QUARANTINE_DIR.mkdir(parents=True, exist_ok=True)
USER_AGENT = {'User-Agent': 'NorthAmericaCityImageResearch/1.0 (city validation and attribution audit)'}
OPENVERSE_URL = 'https://api.openverse.org/v1/images/'
ALLOWED_LICENSES = {'cc0', 'pdm', 'by', 'by-sa'}
MIN_IMAGE_PIXELS = 400_000
COUNTRY_CODE = {
    'Antigua and Barbuda': 'ag', 'Bahamas': 'bs', 'Barbados': 'bb', 'Belize': 'bz', 'Canada': 'ca',
    'Costa Rica': 'cr', 'Cuba': 'cu', 'Dominica': 'dm', 'Dominican Republic': 'do', 'El Salvador': 'sv',
    'Grenada': 'gd', 'Guatemala': 'gt', 'Haiti': 'ht', 'Honduras': 'hn', 'Jamaica': 'jm',
    'Mexico': 'mx', 'Nicaragua': 'ni', 'Panama': 'pa', 'Saint Kitts and Nevis': 'kn', 'Saint Lucia': 'lc',
    'Saint Vincent and the Grenadines': 'vc', 'Trinidad and Tobago': 'tt', 'United States of America': 'us',
}


def slugify(value):
    text = str(value or '').strip().lower().replace("'", '')
    text = text.replace('&', ' and ')
    text = re.sub(r'[^a-z0-9]+', '-', text)
    return text.strip('-')


def normalize_city_name(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii')
    text = text.replace('&', ' and ')
    text = text.replace("'", '').replace('.', ' ')
    text = re.sub(r'[^a-z0-9]+', ' ', text.lower())
    return ' '.join(text.split())


def build_search_variants(city, country):
    ascii_city = normalize_city_name(city)
    country_search = 'USA' if country == 'United States of America' else country
    queries = [
        f'"{city}" "{country}"',
        f'"{city}" "{country_search}"',
        f'"{ascii_city}" "{country_search}"',
        f'"{city}" "{country_search}" skyline',
    ]
    if city == 'Andros Town' and country == 'Bahamas':
        queries.append('"Andros" "Bahamas"')
    if city == 'Rock Sound' and country == 'Bahamas':
        queries.append('"Rock Sound" "Eleuthera"')
    if city == 'West Palm Beach' and country == 'United States of America':
        queries.extend(['"West Palm Beach" Florida', '"West Palm Beach" Florida skyline'])
    return list(dict.fromkeys(queries))


def score_candidate_match(city, country, text):
    normalized = normalize_city_name(text)
    city_key = normalize_city_name(city)
    country_key = normalize_city_name(country)
    compact_normalized = normalized.replace(' ', '')
    compact_city_key = city_key.replace(' ', '')
    if city_key not in normalized and compact_city_key not in compact_normalized:
        return 0
    disqualifying_terms = [
        'portrait', 'selfie', 'panda', 'kangaroo', 'bird', 'moth', 'insect', 'flower',
        'plant specimen', 'animal', 'coin', 'medal', 'stamp', 'coat of arms', 'crest',
        'map', 'satellite', 'illustration', 'logo', 'telegraph', 'document', 'license plate',
        'motor buggy', 'inside the plane', 'airplane cabin', 'flood irrigation',
    ]
    if any(term in normalized for term in disqualifying_terms):
        return 0
    score = 0
    if city_key in normalized:
        score += 60
    if country_key in normalized:
        score += 25
    if city_key.replace(' ', '') in normalized.replace(' ', ''):
        score += 10
    for phrase in [city_key, f'{city_key} {country_key}', f'{city_key} city', f'{city_key} skyline', country_key]:
        if phrase in normalized:
            score += 15
    if any(token in normalized for token in ['skyline', 'downtown', 'city hall', 'harbour', 'harbor', 'centre', 'center']):
        score += 5
    if any(token in normalized for token in ['stock', 'nature', 'forest', 'beach', 'mountain', 'generic']):
        score -= 40
    return score


def make_output_filename(country, city):
    code = COUNTRY_CODE.get(country, country[:2].lower())
    return f'{code}-{slugify(city)}.webp'


def load_inventory():
    wb = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    ws = wb['City_Details']
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx_country = header.index('country')
    idx_city = header.index('City')
    records = []
    seen = set()
    for row in rows[1:]:
        if not row or not row[idx_country] or not row[idx_city]:
            continue
        country = str(row[idx_country]).strip()
        city = str(row[idx_city]).strip()
        if not country or not city:
            continue
        key = (country, city)
        if key in seen:
            continue
        seen.add(key)
        records.append({
            'country': country,
            'city': city,
            'normalized_city': normalize_city_name(city),
            'filename': None,
        })
    return records


def load_manifest():
    if not MANIFEST_PATH.exists():
        return {'continent': 'North America', 'source_workbook': 'Data/North America.xlsx', 'target_folder': 'backend/public/Images/North America cities images', 'cities': []}
    try:
        manifest = json.loads(MANIFEST_PATH.read_text(encoding='utf-8'))
        return manifest if isinstance(manifest, dict) else {'continent': 'North America', 'source_workbook': 'Data/North America.xlsx', 'target_folder': 'backend/public/Images/North America cities images', 'cities': []}
    except (OSError, json.JSONDecodeError):
        return {'continent': 'North America', 'source_workbook': 'Data/North America.xlsx', 'target_folder': 'backend/public/Images/North America cities images', 'cities': []}


def save_manifest(manifest):
    manifest['updated_at'] = time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())
    temp = MANIFEST_PATH.with_name(MANIFEST_PATH.name + '.tmp')
    temp.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    temp.replace(MANIFEST_PATH)


def fetch_json(url):
    request = urllib.request.Request(url, headers=USER_AGENT)
    for attempt in range(1, 7):
        try:
            with urllib.request.urlopen(request, timeout=45) as resp:
                return json.loads(resp.read().decode('utf-8'))
        except urllib.error.HTTPError as exc:
            if exc.code == 429:
                retry_after = exc.headers.get('Retry-After')
                seconds = int(retry_after) if retry_after and retry_after.isdigit() else 8
                if attempt < 6:
                    time.sleep(seconds + attempt)
                    continue
                raise RuntimeError(f'Openverse rate-limit reached: Retry-After={retry_after}') from exc
            if exc.code in {403, 500, 502, 503, 504} and attempt < 6:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError(f'HTTP {exc.code}: {exc.reason}') from exc
        except Exception as exc:
            if attempt < 6:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError(f'network_error:{exc}') from exc
    raise RuntimeError('unreachable_fetch')


def fetch_bytes(url):
    request = urllib.request.Request(url, headers=USER_AGENT)
    for attempt in range(1, 7):
        try:
            with urllib.request.urlopen(request, timeout=45) as resp:
                return resp.read()
        except urllib.error.HTTPError as exc:
            if exc.code == 429:
                retry_after = exc.headers.get('Retry-After')
                seconds = int(retry_after) if retry_after and retry_after.isdigit() else 8
                if attempt < 6:
                    time.sleep(seconds + attempt)
                    continue
                raise RuntimeError(f'429 rate limit on image URL: Retry-After={retry_after}') from exc
            if exc.code in {403, 500, 502, 503, 504} and attempt < 6:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError(f'HTTP {exc.code}: {exc.reason} for {url}') from exc
        except Exception as exc:
            if attempt < 6:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError(f'download_error:{exc}') from exc
    raise RuntimeError(f'failed_to_fetch:{url}')


def validate_image_bytes(data, city, country):
    try:
        image = Image.open(BytesIO(data))
        image = ImageOps.exif_transpose(image)
        if image.mode in {'RGBA', 'LA', 'P'}:
            image = image.convert('RGB')
        image.load()
        if image.width < 400 or image.height < 400:
            raise ValueError(f'image too small: {image.size}')
        if image.width * image.height < MIN_IMAGE_PIXELS:
            raise ValueError(f'image below minimum pixel threshold: {image.size}')
        return image
    except Exception as exc:
        raise RuntimeError(f'Invalid image for {country} / {city}: {exc}') from exc


def build_candidate_from_openverse(item):
    if not item.get('url') or not item.get('foreign_landing_url'):
        return None
    license_name = (item.get('license') or '').lower()
    if license_name not in ALLOWED_LICENSES:
        return None
    width = int(item.get('width') or 0)
    height = int(item.get('height') or 0)
    if width < 400 or height < 400 or width * height < MIN_IMAGE_PIXELS:
        return None
    return {
        'source': (item.get('source') or 'openverse').lower(),
        'url': item.get('url'),
        'landing_url': item.get('foreign_landing_url'),
        'title': item.get('title') or '',
        'creator': item.get('creator') or 'Unknown creator',
        'license': license_name,
        'license_version': item.get('license_version') or '',
        'license_url': item.get('license_url') or '',
        'attribution': item.get('attribution') or '',
        'width': width,
        'height': height,
    }


def find_openverse_candidates(city, country):
    results = {}
    for query in build_search_variants(city, country):
        params = urllib.parse.urlencode({
            'q': query,
            'page': 1,
            'page_size': 10,
            'license': ','.join(sorted(ALLOWED_LICENSES)),
            'mature': 'false',
        })
        try:
            data = fetch_json(f'{OPENVERSE_URL}?{params}')
        except RuntimeError:
            continue
        for item in data.get('results', []):
            built = build_candidate_from_openverse(item)
            if not built:
                continue
            match_score = score_candidate_match(city, country, json.dumps({
                'title': item.get('title', ''),
                'source': item.get('source', ''),
                'attribution': item.get('attribution', ''),
                'tags': item.get('tags', []),
            }, ensure_ascii=False))
            if match_score <= 0:
                continue
            key = built['url']
            if key not in results or match_score > results[key]['score']:
                built['score'] = match_score
                results[key] = built
    return sorted(results.values(), key=lambda item: item['score'], reverse=True)[:8]


def save_image(country, city, candidate, replace_existing=False):
    filename = make_output_filename(country, city)
    target = OUTDIR / filename
    if target.exists() and not replace_existing:
        try:
            with Image.open(target) as existing:
                existing.load()
                if existing.format == 'WEBP' and existing.width >= 400 and existing.height >= 400 and existing.width * existing.height >= MIN_IMAGE_PIXELS:
                    return {'status': 'already_exists', 'filename': filename, 'file': str(target), 'dimensions': [existing.width, existing.height]}
        except Exception:
            pass
    payload = fetch_bytes(candidate['url'])
    if not payload or len(payload) < 128:
        raise RuntimeError(f'Empty or suspicious payload downloaded for {country} / {city}')
    image = validate_image_bytes(payload, city, country)
    if image.width > 1800 or image.height > 1800:
        image.thumbnail((1800, 1800), Image.Resampling.LANCZOS)
    temp = target.with_suffix('.webp.part')
    image.save(temp, 'WEBP', quality=85, optimize=True)
    temp.replace(target)
    with Image.open(target) as saved:
        saved.load()
        if saved.width < 400 or saved.height < 400 or saved.width * saved.height < MIN_IMAGE_PIXELS:
            raise RuntimeError(f'Validated file too small after conversion: {saved.size}')
    return {'status': 'downloaded', 'filename': filename, 'file': str(target), 'dimensions': [saved.width, saved.height]}


def should_skip_reprocessing(entry):
    if not isinstance(entry, dict):
        return False
    validation = entry.get('validation') or {}
    decision = validation.get('visual_review')
    if decision in {'approved', 'rejected', 'needs_review'}:
        return True
    status = entry.get('download_status')
    if status in {'final_approved', 'visual_rejected', 'downloaded_pending_review', 'existing_review_required'}:
        return True
    if status == 'unresolved' and not entry.get('filename'):
        return True
    return False


def is_final_approval_ready(entry):
    filename = entry.get('filename')
    if not filename:
        return False
    target = OUTDIR / filename
    if not target.exists():
        return False
    try:
        with Image.open(target) as image:
            image.load()
            width, height = image.size
            if image.format != 'WEBP':
                return False
            if width < 400 or height < 400 or width * height < MIN_IMAGE_PIXELS:
                return False
    except Exception:
        return False
    required_fields = ['source_url', 'source_page', 'creator', 'license', 'license_url', 'attribution']
    if not all(entry.get(field) for field in required_fields):
        return False
    if not str(entry.get('creator', '')).strip() or str(entry.get('creator', '')).strip().lower() == 'unknown creator':
        return False
    license_name = re.sub(r'^cc\s+', '', str(entry.get('license', '')).strip().lower())
    license_name = license_name.split()[0] if license_name else ''
    if license_name not in ALLOWED_LICENSES:
        return False
    validation = entry.get('validation', {})
    if validation.get('visual_review') != 'approved':
        return False
    evidence = validation.get('visual_review_evidence', {})
    if not all(evidence.get(field) for field in ['reviewer', 'reviewed_at', 'contact_sheet', 'note']):
        return False
    city_relevance = validation.get('city_relevance', '')
    if not city_relevance or 'pending' in city_relevance.lower():
        return False
    return True


def record_visual_review(entry, decision, reviewer, contact_sheet, note):
    if decision not in {'approved', 'rejected', 'needs_review'}:
        raise ValueError(f'Unsupported visual-review decision: {decision}')
    if not all(str(value or '').strip() for value in [reviewer, contact_sheet, note]):
        raise ValueError('Reviewer, contact sheet, and review note are required.')

    validation = entry.setdefault('validation', {})
    validation['visual_review'] = decision
    validation['visual_review_evidence'] = {
        'reviewer': reviewer.strip(),
        'reviewed_at': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
        'contact_sheet': contact_sheet.strip(),
        'note': note.strip(),
    }
    validation['city_relevance'] = note.strip()
    entry.setdefault('review_history', []).append({
        'decision': decision,
        **validation['visual_review_evidence'],
        'filename': entry.get('filename'),
        'note': note.strip(),
    })

    if decision == 'rejected':
        entry['download_status'] = 'visual_rejected'
    elif is_final_approval_ready(entry):
        entry['download_status'] = 'final_approved'
        validation['approval_status'] = 'final_approved'
        validation['final_reason'] = 'Technical, attribution, license, and documented visual-review checks passed.'
    else:
        entry['download_status'] = 'existing_review_required' if entry.get('filename') else 'unresolved'
        validation['approval_status'] = 'review_required'
    return entry


def quarantine_visual_rejection(entry, reviewer, contact_sheet, note):
    filename = entry.get('filename')
    source = OUTDIR / filename if filename else None
    if not source or not source.is_file():
        raise FileNotFoundError(f'Cannot quarantine missing image for {entry.get("city")}')

    record_visual_review(entry, 'rejected', reviewer, contact_sheet, note)
    destination = QUARANTINE_DIR / f'{source.stem}.visual-rejected{source.suffix}'
    suffix = 2
    while destination.exists():
        destination = QUARANTINE_DIR / f'{source.stem}.visual-rejected-{suffix}{source.suffix}'
        suffix += 1
    shutil.move(str(source), str(destination))
    entry['rejected_filename'] = filename
    entry['rejected_candidate_file'] = destination.relative_to(ROOT).as_posix()
    entry['filename'] = None
    entry['download_status'] = 'visual_rejected'
    entry['unresolved_reason'] = note.strip()
    return entry


def finalize_existing_approvals(manifest):
    approved = 0
    for entry in manifest['cities']:
        if is_final_approval_ready(entry):
            entry['download_status'] = 'final_approved'
            entry['validation']['approval_status'] = 'final_approved'
            entry['validation']['final_reason'] = 'File is valid, mapped to the correct city, and has licensing metadata.'
            approved += 1
        elif entry.get('download_status') == 'final_approved':
            filename = entry.get('filename')
            target = OUTDIR / filename if filename else None
            entry['download_status'] = 'existing_review_required' if target and target.exists() else 'unresolved'
            entry.setdefault('validation', {})['approval_status'] = 'review_required'
            entry['validation']['final_reason'] = 'Final approval requires explicit visual-review evidence and complete attribution.'
    return approved


def seed_manifest():
    manifest = load_manifest()
    manifest.setdefault('continent', 'North America')
    manifest.setdefault('source_workbook', 'Data/North America.xlsx')
    manifest.setdefault('target_folder', 'backend/public/Images/North America cities images')
    manifest.setdefault('allowed_licenses', sorted(ALLOWED_LICENSES))
    manifest.setdefault('cities', [])
    by_key = {(entry.get('country'), entry.get('city')): entry for entry in manifest['cities']}
    ordered = []
    for record in load_inventory():
        key = (record['country'], record['city'])
        entry = by_key.get(key, {
            'country': record['country'],
            'city': record['city'],
            'filename': None,
            'download_status': 'pending',
            'validation': {},
            'sources_attempted': [],
        })
        entry.setdefault('country', record['country'])
        entry.setdefault('city', record['city'])
        entry.setdefault('filename', None)
        entry.setdefault('download_status', 'pending')
        entry.setdefault('validation', {})
        entry.setdefault('sources_attempted', [])
        expected_name = make_output_filename(record['country'], record['city'])
        target = OUTDIR / expected_name
        if target.exists():
            entry['filename'] = expected_name
            try:
                with Image.open(target) as image:
                    image.load()
                    if image.width >= 400 and image.height >= 400 and image.width * image.height >= MIN_IMAGE_PIXELS:
                        validation = entry.setdefault('validation', {})
                        validation.update({'file_integrity': 'valid WebP', 'dimensions': [image.width, image.height]})
                        validation.setdefault('city_relevance', 'existing file retained for review')
                        if validation.get('visual_review') == 'rejected':
                            entry['download_status'] = 'visual_rejected'
                        elif is_final_approval_ready(entry):
                            entry['download_status'] = 'final_approved'
                            validation['approval_status'] = 'final_approved'
                        else:
                            entry['download_status'] = 'existing_review_required'
            except Exception as exc:
                entry['download_status'] = 'invalid_existing_file'
                entry.setdefault('validation', {})['file_integrity'] = f'failed: {exc}'
        ordered.append(entry)
    manifest['cities'] = ordered
    return manifest


if __name__ == '__main__':
    print('North America city image batch: validating source selection and resuming image downloads.')
    manifest = seed_manifest()
    save_manifest(manifest)
    total = len(manifest['cities'])
    for index, entry in enumerate(manifest['cities'], start=1):
        country = entry['country']
        city = entry['city']
        if should_skip_reprocessing(entry):
            print(f'{index}/{total} skip {country} | {city} -> status={entry.get("download_status")} decision={entry.get("validation", {}).get("visual_review")}')
            continue
        target = OUTDIR / make_output_filename(country, city)
        if target.exists():
            try:
                with Image.open(target) as image:
                    image.load()
                    if image.width >= 400 and image.height >= 400 and image.width * image.height >= MIN_IMAGE_PIXELS:
                        entry['filename'] = target.name
                        validation = entry.setdefault('validation', {})
                        validation.update({'file_integrity': 'valid WebP', 'dimensions': [image.width, image.height]})
                        validation.setdefault('city_relevance', 'existing file retained for review')
                        if validation.get('visual_review') == 'rejected':
                            entry['download_status'] = 'visual_rejected'
                        elif is_final_approval_ready(entry):
                            entry['download_status'] = 'final_approved'
                            validation['approval_status'] = 'final_approved'
                        else:
                            entry['download_status'] = 'existing_review_required'
                        save_manifest(manifest)
                        print(f'{index}/{total} existing {country} | {city} -> {target.name}')
                        continue
            except Exception:
                pass
        try:
            results = find_openverse_candidates(city, country)
            if not results:
                entry['download_status'] = 'unresolved'
                entry['unresolved_reason'] = 'No suitable licensed city image matching the city and country was found in Openverse.'
                entry['validation'] = {'city_relevance': 'not verified'}
                save_manifest(manifest)
                print(f'{index}/{total} unresolved {country} | {city}')
                continue
            for candidate in results:
                try:
                    saved = save_image(country, city, candidate, replace_existing=True)
                    entry['filename'] = saved['filename']
                    entry['source_url'] = candidate['url']
                    entry['source_page'] = candidate['landing_url']
                    entry['creator'] = candidate['creator']
                    entry['license'] = candidate['license']
                    entry['license_version'] = candidate['license_version']
                    entry['license_url'] = candidate['license_url']
                    entry['attribution'] = candidate['attribution'] or f"{candidate['creator']} licensed under {candidate['license']}"
                    entry['download_status'] = 'downloaded_pending_review'
                    entry['validation'] = {'file_integrity': 'valid WebP', 'dimensions': saved['dimensions'], 'city_relevance': 'source title and metadata match the intended city; pending final visual verification'}
                    entry['unresolved_reason'] = None
                    entry.setdefault('sources_attempted', []).append({
                        'source': candidate['source'],
                        'source_url': candidate['url'],
                        'landing_url': candidate['landing_url'],
                        'license': candidate['license'],
                        'status': 'downloaded',
                    })
                    if is_final_approval_ready(entry):
                        entry['download_status'] = 'final_approved'
                        entry['validation']['approval_status'] = 'final_approved'
                    save_manifest(manifest)
                    print(f'{index}/{total} saved {country} | {city} -> {saved["filename"]}')
                    break
                except Exception as exc:
                    entry.setdefault('sources_attempted', []).append({
                        'source': candidate['source'],
                        'source_url': candidate['url'],
                        'landing_url': candidate['landing_url'],
                        'license': candidate['license'],
                        'status': 'download_failed',
                        'reason': str(exc),
                    })
                    continue
            else:
                entry['download_status'] = 'unresolved'
                entry['unresolved_reason'] = 'All candidate Openverse results failed validation or were below the required resolution threshold.'
                save_manifest(manifest)
                print(f'{index}/{total} failed validation {country} | {city}')
        except Exception as exc:
            entry['download_status'] = 'retryable_error'
            entry['unresolved_reason'] = str(exc)
            save_manifest(manifest)
            print(f'{index}/{total} retryable error {country} | {city} -> {exc}')

    approved = finalize_existing_approvals(manifest)
    save_manifest(manifest)
    print(f'Done. processed={total} final_approved={approved}')
    print(f'Manifest: {MANIFEST_PATH}')
