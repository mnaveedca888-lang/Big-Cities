from docx import Document
import re

# Load the document
doc = Document('Data/Azerbaijan/Azerbaijan.docx')

# Extract all text
full_text = []
for para in doc.paragraphs:
    if para.text.strip():
        full_text.append(para.text)

# Also extract from tables
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.strip():
                full_text.append(cell.text)

# Join all text
text_combined = '\n'.join(full_text)

# Define cities
cities = ['GANJA', 'LANKARAN', 'NAKHCHIVAN', 'GABALA', 'LACHIN', 'FUZULI', 'ZANGILAN', 'ZAQATALA']

# Save the full text to file for inspection
with open('tmp_full_docx_content.txt', 'w', encoding='utf-8') as f:
    f.write(text_combined)

print(f"Full document content saved to tmp_full_docx_content.txt ({len(text_combined)} characters)")
print(f"\nSearching for cities: {', '.join(cities)}")

# Find city mentions
for city in cities:
    count = len(re.findall(rf'\b{city}\b', text_combined, re.IGNORECASE))
    print(f"{city}: found {count} times")
